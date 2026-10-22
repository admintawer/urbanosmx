# -*- coding: utf-8 -*-

from odoo import models, fields, _
from odoo.exceptions import ValidationError
import base64
from io import BytesIO

from openpyxl import load_workbook

class ImportBlanket(models.TransientModel):
    _name = 'import.blanket.wizard'
    _description = 'Import vendors agreement offers'

    subtype = fields.Selection(string="Criterio", selection=[('time','Tiempo de entrega'),('price','Mejor precio'),('time_price', 'Tiempo + Precio')])
    requisition_id = fields.Many2one('purchase.requisition')
    company_id = fields.Many2one('res.company', related='requisition_id.company_id', readonly=True)
    vendor_domain = fields.Char()
    vendor_ids = fields.Many2many('res.partner', string="Proveedores", domain="['|', ('company_id', '=', False), ('company_id', '=', company_id)]")
    vendor_id = fields.Many2one('res.partner', string='Proveedor')
    xlsx_file = fields.Binary('File')

    def import_bl_xlsx(self):
        blanket_lines = self.env['pr.blanket.lines'].search([('requisition_id','=',self.requisition_id.id),('partner_id','=', self.vendor_id.id)])
        if blanket_lines:
            raise ValidationError('Ya existe una oferta importada para este proveedor, revisa por favor.')
        else:
            ncounter = 1
            skipped_line_no = {}
            workbook = load_workbook(BytesIO(base64.b64decode(self.xlsx_file)), data_only=True)
            sheet = workbook.active
            for row in sheet.iter_rows(min_row=2, values_only=True):
                blanket_vals = {}
                line_error = False
                if row[0] not in (None, ""):
                    product = self.env['product.product'].browse(int(row[0])).exists()
                    if product:
                        blanket_vals.update({
                            'requisition_id': self.requisition_id.id,
                            'partner_id': self.vendor_id.id,
                            'product_id': product.id,
                        })
                        if row[4] not in (None, ""):
                            price_unit = row[4]
                            blanket_vals.update({
                                'price_unit': price_unit
                            })
                        else:
                            line_error = True
                            skipped_line_no[str(ncounter)] = " - Hay un error en el precio unitario, por favor revisar"
                        if self.subtype in ('time','time_price'):
                            if row[5] not in (None, ""):
                                schedule_date = fields.Date.to_string(row[5])
                                blanket_vals.update({
                                    'schedule_date': schedule_date,
                                })
                            else:
                                line_error = True
                                skipped_line_no[str(ncounter)] = " - Hay un error en la fecha de entrega, por favor revisar"
                        if not line_error:
                            self.env['pr.blanket.lines'].create(blanket_vals)
                    else:
                        line_error = True
                        skipped_line_no[str(ncounter)] = " - El codigo del producto no se encuentra, por favor revisar"
                        
                else:
                    line_error = True
                    skipped_line_no[str(ncounter)] = " - El codigo del producto esta vacio. "
                ncounter += 1
            if skipped_line_no:
                raise ValidationError(str(skipped_line_no))

                    
