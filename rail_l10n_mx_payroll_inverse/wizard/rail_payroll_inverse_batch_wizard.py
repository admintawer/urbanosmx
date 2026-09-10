# -*- coding: utf-8 -*-

import base64
import io

try:
    from openpyxl import load_workbook
except ImportError:  # pragma: no cover - handled at runtime inside Odoo
    load_workbook = None

from odoo import Command, fields, models, _
from odoo.exceptions import UserError


class RailPayrollInverseBatchWizard(models.TransientModel):
    _name = 'rail.payroll.inverse.batch.wizard'
    _description = 'Importación masiva de cálculo inverso MX'

    _MONTH_SELECTION = [
        ('01', 'Enero / Periodo 1'), ('02', 'Febrero / Periodo 2'),
        ('03', 'Marzo / Periodo 3'), ('04', 'Abril / Periodo 4'),
        ('05', 'Mayo / Periodo 5'), ('06', 'Junio / Periodo 6'),
        ('07', 'Julio / Periodo 7'), ('08', 'Agosto / Periodo 8'),
        ('09', 'Septiembre / Periodo 9'), ('10', 'Octubre / Periodo 10'),
        ('11', 'Noviembre / Periodo 11'), ('12', 'Diciembre / Periodo 12'),
    ]

    company_id = fields.Many2one('res.company', string='Compañía', required=True, default=lambda self: self.env.company)
    currency_id = fields.Many2one(related='company_id.currency_id', readonly=True)
    file = fields.Binary(string='Archivo XLSX', required=True)
    filename = fields.Char(string='Nombre archivo')
    date_from = fields.Date(string='Desde', required=True, default=lambda self: fields.Date.today().replace(day=1))
    date_to = fields.Date(string='Hasta', required=True, default=fields.Date.today)
    effective_date = fields.Date(string='Fecha efectiva por defecto', required=True, default=fields.Date.today)
    tolerance = fields.Monetary(string='Tolerancia', default=0.05, currency_field='currency_id')
    max_iterations = fields.Integer(string='Iteraciones máximas', default=40)
    auto_apply = fields.Boolean(string='Crear versiones al calcular')
    line_ids = fields.One2many('rail.payroll.inverse.batch.line', 'wizard_id', string='Líneas')
    state = fields.Selection([
        ('draft', 'Borrador'),
        ('loaded', 'Cargado'),
        ('calculated', 'Calculado'),
        ('applied', 'Aplicado'),
    ], default='draft', string='Estado', readonly=True)

    def _normalize_header(self, value):
        return (str(value or '').strip().lower()
                .replace('á', 'a').replace('é', 'e').replace('í', 'i')
                .replace('ó', 'o').replace('ú', 'u').replace('ñ', 'n'))

    def _get_header_map(self, row):
        aliases = {
            'employee_number': {'no. empleado', 'no empleado', 'numero empleado', 'num empleado'},
            'candidate_name': {'nombre candidato', 'candidato', 'nombre empleado', 'empleado'},
            'target_net_amount': {'monto neto', 'neto deseado', 'neto objetivo', 'neto', 'monto'},
            'payroll_schedule_pay': {
                'periodicidad nomina', 'periodicidad de nomina', 'periodo nomina', 'tipo nomina',
            },
            'target_schedule_pay': {
                'periodicidad neto', 'periodicidad del neto', 'periodicidad',
                'tipo sueldo', 'tipo de sueldo', 'periodo neto',
            },
            'effective_date': {'fecha efectiva', 'fecha version', 'fecha'},
            'fixed_sbc': {'sbc fijo', 'sbc', 'sbc manual'},
            'ultima_nomina': {'ultima nomina del mes', 'ultima nomina', 'ultimo periodo del mes'},
            'mes': {'mes', 'mes nomina', 'mes de la nomina'},
            'structure': {'estructura salarial', 'estructura', 'codigo estructura', 'estructura a simular'},
            'contract_start_date': {'fecha estimada ingreso', 'fecha de ingreso', 'fecha ingreso', 'inicio contrato'},
            'holiday_bonus_rate': {'prima vacacional', 'prima vacacional %', 'prima vacacional porcentaje'},
        }
        header_map = {}
        for index, cell in enumerate(row, start=1):
            header = self._normalize_header(cell.value)
            for key, names in aliases.items():
                if header in names:
                    header_map[key] = index
        if 'target_net_amount' not in header_map:
            raise UserError(_('El archivo debe incluir la columna Neto deseado.'))
        if 'employee_number' not in header_map and 'candidate_name' not in header_map:
            raise UserError(_('El archivo debe incluir No. empleado o Nombre candidato.'))
        return header_map

    def _parse_target_schedule_pay(self, value):
        normalized = self._normalize_header(value).replace('-', ' ').replace('_', ' ')
        aliases = {
            'daily': {'daily', 'diario', 'dia', '1 dia'},
            'weekly': {'weekly', 'semanal', 'semana', '7 dias'},
            '10_days': {'10 days', '10 dias', 'cada 10 dias', 'decenal'},
            '14_days': {'14 days', '14 dias', 'cada 14 dias'},
            'bi-weekly': {'bi weekly', 'biweekly', 'quincenal', 'quincena', '15 dias'},
            'monthly': {'monthly', 'mensual', 'mes', '30 dias'},
            'bi-monthly': {'bi monthly', 'bimonthly', 'bimestral', '2 meses', '60 dias'},
        }
        if not normalized:
            return False
        for key, names in aliases.items():
            if normalized in names:
                return key
        raise UserError(_('Periodicidad de neto no reconocida: %s') % value)

    def _find_structure(self, value):
        if value in (None, False, ''):
            return self.env['hr.payroll.structure']
        text = str(value).strip()
        structures = self.env['hr.payroll.structure'].search([
            '|', ('code', '=', text), ('name', '=', text),
        ], limit=2)
        if not structures:
            return structures
        if len(structures) > 1:
            raise UserError(_('La estructura salarial %s coincide con más de un registro. Use el código único de la estructura.') % text)
        return structures

    def _parse_boolean(self, value):
        if value in (None, False, ''):
            return False
        if value is True:
            return True
        normalized = self._normalize_header(value)
        if normalized in {'1', 'si', 'sí', 'true', 'verdadero', 'x'}:
            return True
        if normalized in {'0', 'no', 'false', 'falso'}:
            return False
        raise UserError(_('Valor de Última nómina del mes no reconocido: %s') % value)

    def _parse_month(self, value):
        if value in (None, False, ''):
            return False
        try:
            month = int(float(value))
        except (TypeError, ValueError):
            normalized = self._normalize_header(value)
            names = {
                'enero': 1, 'febrero': 2, 'marzo': 3, 'abril': 4, 'mayo': 5, 'junio': 6,
                'julio': 7, 'agosto': 8, 'septiembre': 9, 'octubre': 10, 'noviembre': 11, 'diciembre': 12,
            }
            month = names.get(normalized, 0)
        if not 1 <= month <= 12:
            raise UserError(_('Mes de nómina no reconocido: %s') % value)
        return '%02d' % month

    def _rail_reopen_wizard_action(self):
        self.ensure_one()
        return {
            'type': 'ir.actions.act_window',
            'name': _('Importar cálculo inverso'),
            'res_model': self._name,
            'view_mode': 'form',
            'res_id': self.id,
            'target': 'new',
            'context': dict(self.env.context),
        }

    def action_load_file(self):
        self.ensure_one()
        if load_workbook is None:
            raise UserError(_('No está disponible la librería openpyxl para leer archivos XLSX.'))
        if not self.file:
            raise UserError(_('Seleccione un archivo XLSX.'))
        workbook = load_workbook(io.BytesIO(base64.b64decode(self.file)), data_only=True)
        sheet = workbook.active
        header_map = self._get_header_map(next(sheet.iter_rows(min_row=1, max_row=1)))
        commands = [Command.clear()]
        Employee = self.env['hr.employee'].with_context(active_test=False)
        for row_number, row in enumerate(sheet.iter_rows(min_row=2), start=2):
            employee_number = (
                row[header_map['employee_number'] - 1].value
                if header_map.get('employee_number') else False
            )
            candidate_name = (
                row[header_map['candidate_name'] - 1].value
                if header_map.get('candidate_name') else False
            )
            target = row[header_map['target_net_amount'] - 1].value
            if not employee_number and not candidate_name and not target:
                continue
            employee = self.env['hr.employee']
            lookup_message = False
            if employee_number:
                employee = Employee.search([
                    ('registration_number', '=', str(employee_number).strip())
                ], limit=2)
                if len(employee) > 1:
                    lookup_message = _('Más de un empleado coincide con el número indicado.')
                    employee = self.env['hr.employee']
            elif candidate_name:
                employee = Employee.search([('name', '=', str(candidate_name).strip())], limit=2)
                if len(employee) > 1:
                    lookup_message = _(
                        'Más de un empleado coincide exactamente con el nombre. Use No. empleado para identificarlo.'
                    )
                    employee = self.env['hr.employee']
            effective = self.effective_date
            if header_map.get('effective_date'):
                effective = row[header_map['effective_date'] - 1].value or effective
            fixed_sbc = 0.0
            if header_map.get('fixed_sbc'):
                fixed_sbc = row[header_map['fixed_sbc'] - 1].value or 0.0
            target_schedule_pay = False
            if header_map.get('target_schedule_pay'):
                target_schedule_pay = self._parse_target_schedule_pay(
                    row[header_map['target_schedule_pay'] - 1].value
                )
            payroll_schedule_pay = False
            if header_map.get('payroll_schedule_pay'):
                payroll_schedule_pay = self._parse_target_schedule_pay(
                    row[header_map['payroll_schedule_pay'] - 1].value
                )
            contract_start_date = False
            if header_map.get('contract_start_date'):
                contract_start_date = row[header_map['contract_start_date'] - 1].value or False
            holiday_bonus_rate = 0.0
            if header_map.get('holiday_bonus_rate'):
                holiday_bonus_rate = row[header_map['holiday_bonus_rate'] - 1].value or 0.0
            ultima_nomina = False
            if header_map.get('ultima_nomina'):
                ultima_nomina = self._parse_boolean(row[header_map['ultima_nomina'] - 1].value)
            mes = False
            if header_map.get('mes'):
                mes = self._parse_month(row[header_map['mes'] - 1].value)
            structure = self.env['hr.payroll.structure']
            structure_value = False
            if header_map.get('structure'):
                structure_value = row[header_map['structure'] - 1].value
                structure = self._find_structure(structure_value)
            if ultima_nomina and not mes:
                raise UserError(_('Fila %s: indique el Mes cuando marca Última nómina del mes.') % row_number)
            resolved_candidate_name = str(candidate_name or '').strip() or (employee.name if employee else '')
            status = 'loaded'
            message = False
            if lookup_message:
                status = 'error'
                message = lookup_message
            elif structure_value and not structure:
                status = 'error'
                message = _('Estructura salarial no encontrada: %s') % structure_value
            elif not employee and not resolved_candidate_name:
                status = 'error'
                message = _('Indique No. empleado o Nombre candidato.')
            commands.append(Command.create({
                'row_number': row_number,
                'employee_number': str(employee_number or '').strip(),
                'candidate_name': resolved_candidate_name,
                'employee_id': employee.id,
                'target_net_amount': float(target or 0.0),
                'payroll_schedule_pay': payroll_schedule_pay,
                'target_schedule_pay': target_schedule_pay,
                'ultima_nomina': ultima_nomina,
                'mes': mes,
                'effective_date': effective,
                'contract_start_date': contract_start_date,
                'holiday_bonus_rate': float(holiday_bonus_rate or 0.0),
                'fixed_sbc': float(fixed_sbc or 0.0),
                'use_fixed_sbc': bool(fixed_sbc),
                'structure_id': structure.id,
                'status': status,
                'message': message,
            }))
        self.write({'line_ids': commands, 'state': 'loaded'})
        return self._rail_reopen_wizard_action()

    def action_calculate(self):
        self.ensure_one()
        for line in self.line_ids.filtered(lambda item: item.status != 'error'):
            line.action_calculate()
            if self.auto_apply and line.status == 'calculated' and line.employee_id:
                line.action_apply()
        self.state = (
            'applied'
            if self.auto_apply and self.line_ids and all(line.status == 'applied' for line in self.line_ids)
            else 'calculated'
        )
        return self._rail_reopen_wizard_action()

    def action_apply(self):
        self.ensure_one()
        for line in self.line_ids.filtered(lambda item: item.status == 'calculated' and item.employee_id):
            line.action_apply()
        standalone_candidates = self.line_ids.filtered(
            lambda item: item.status == 'calculated' and not item.employee_id
        )
        for line in standalone_candidates:
            if not line.message:
                line.message = _(
                    'Simulación de candidato calculada. No se crea una versión laboral hasta que exista una ficha de empleado.'
                )
        self.state = (
            'applied'
            if self.line_ids and all(
                line.status == 'applied' or (line.status == 'calculated' and not line.employee_id)
                for line in self.line_ids.filtered(lambda item: item.status != 'error')
            )
            else 'calculated'
        )
        return self._rail_reopen_wizard_action()


class RailPayrollInverseBatchLine(models.TransientModel):
    _name = 'rail.payroll.inverse.batch.line'
    _description = 'Línea importación cálculo inverso MX'

    wizard_id = fields.Many2one('rail.payroll.inverse.batch.wizard', required=True, ondelete='cascade')
    row_number = fields.Integer(string='Fila')
    employee_number = fields.Char(string='No. empleado')
    candidate_name = fields.Char(string='Nombre candidato')
    employee_id = fields.Many2one('hr.employee', string='Empleado')
    structure_id = fields.Many2one('hr.payroll.structure', string='Estructura salarial')
    target_net_amount = fields.Monetary(string='Neto objetivo', currency_field='currency_id')
    payroll_schedule_pay = fields.Selection(
        selection=lambda self: self.env['hr.payroll.structure.type']._get_selection_schedule_pay(),
        string='Periodicidad nómina',
    )
    target_schedule_pay = fields.Selection(
        selection=lambda self: self.env['hr.payroll.structure.type']._get_selection_schedule_pay(),
        string='Periodicidad neto',
    )
    ultima_nomina = fields.Boolean(string='Última nómina del mes')
    mes = fields.Selection(RailPayrollInverseBatchWizard._MONTH_SELECTION, string='Mes de la nómina')
    effective_date = fields.Date(string='Fecha efectiva')
    contract_start_date = fields.Date(string='Fecha estimada ingreso')
    holiday_bonus_rate = fields.Float(string='Prima vacacional (%)')
    use_fixed_sbc = fields.Boolean(string='Actualizar SBC')
    fixed_sbc = fields.Float(string='SBC a aplicar')
    currency_id = fields.Many2one(related='wizard_id.currency_id', readonly=True)
    calculated_wage = fields.Monetary(string='Salario calculado', currency_field='currency_id', readonly=True)
    calculated_net = fields.Monetary(string='Neto simulado', currency_field='currency_id', readonly=True)
    result_version_id = fields.Many2one('hr.version', string='Nueva versión', readonly=True)
    status = fields.Selection([
        ('loaded', 'Cargado'),
        ('calculated', 'Calculado'),
        ('applied', 'Aplicado'),
        ('error', 'Error'),
    ], default='loaded', string='Estado')
    message = fields.Text(string='Mensaje')

    def _run_with_individual_wizard(self):
        self.ensure_one()
        employee = self.employee_id
        reference_version = employee._get_version(self.wizard_id.date_from) if employee else self.env['hr.version']
        version = (
            reference_version
            if reference_version and reference_version.contract_date_start and reference_version.schedule_pay
            else self.env['hr.version']
        )
        structure = self.structure_id
        if not structure and reference_version:
            structure = reference_version.structure_type_id.default_struct_id
        if not structure:
            structure = self.env.ref(
                'l10n_mx_hr_payroll.l10n_mx_regular_pay', raise_if_not_found=False
            )
        if not structure:
            raise UserError(_('No fue posible determinar la estructura salarial a simular.'))

        payroll_schedule = (
            version.schedule_pay
            if version
            else self.payroll_schedule_pay or (reference_version.schedule_pay if reference_version else False)
        )
        if not payroll_schedule:
            raise UserError(_(
                'Indique Periodicidad nómina para candidatos o empleados que todavía no tienen versión laboral.'
            ))
        candidate_name = self.candidate_name or (employee.name if employee else False)
        if not employee and not candidate_name:
            raise UserError(_('Indique el nombre del candidato.'))
        if self.ultima_nomina and not version:
            raise UserError(_('Última nómina del mes solo aplica a empleados con versión laboral y nóminas previas.'))

        inverse_wizard = self.env['rail.payroll.inverse.wizard'].create({
            'employee_id': employee.id,
            'candidate_name': candidate_name,
            'version_id': version.id,
            'company_id': self.wizard_id.company_id.id,
            'date_from': self.wizard_id.date_from,
            'date_to': self.wizard_id.date_to,
            'effective_date': self.effective_date or self.wizard_id.effective_date,
            'struct_id': structure.id,
            'simulation_schedule_pay': payroll_schedule,
            'simulation_contract_start_date': (
                self.contract_start_date
                or (version.contract_date_start if version else False)
                or self.effective_date
                or self.wizard_id.effective_date
            ),
            'simulation_holiday_bonus_rate': (
                version.l10n_mx_holiday_bonus_rate
                if version
                else self.holiday_bonus_rate or (reference_version.l10n_mx_holiday_bonus_rate if reference_version else 0.0)
            ),
            'target_net_amount': self.target_net_amount,
            'target_schedule_pay': self.target_schedule_pay or payroll_schedule,
            'ultima_nomina': self.ultima_nomina,
            'mes': self.mes,
            'tolerance': self.wizard_id.tolerance,
            'max_iterations': self.wizard_id.max_iterations,
            'use_fixed_sbc': self.use_fixed_sbc,
            'fixed_sbc': self.fixed_sbc,
        })
        inverse_wizard.action_calculate()
        return inverse_wizard

    def action_calculate(self):
        for line in self:
            try:
                inverse_wizard = line._run_with_individual_wizard()
                line.write({
                    'calculated_wage': inverse_wizard.calculated_wage,
                    'calculated_net': inverse_wizard.calculated_net,
                    'status': 'calculated',
                    'message': inverse_wizard.simulation_log,
                })
            except Exception as error:  # noqa: B902 - show functional error per line
                line.write({'status': 'error', 'message': str(error)})
        return True

    def action_apply(self):
        for line in self:
            try:
                inverse_wizard = line._run_with_individual_wizard()
                inverse_wizard.action_apply_new_version()
                line.write({
                    'calculated_wage': inverse_wizard.calculated_wage,
                    'calculated_net': inverse_wizard.calculated_net,
                    'result_version_id': inverse_wizard.result_version_id.id,
                    'status': 'applied',
                    'message': _('Nueva versión creada: %s') % inverse_wizard.result_version_id.display_name,
                })
            except Exception as error:  # noqa: B902
                line.write({'status': 'error', 'message': str(error)})
        return True
